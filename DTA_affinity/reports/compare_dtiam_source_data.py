"""Compare existing results with the unmodified DTIAM Figure 3 source workbook."""

import hashlib
import json
import statistics
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
ARCHIVE = PROJECT / "data/raw/dtiam_source_data_2025.zip"
MEMBER = "Source data/Source data for Fig.3.xlsx"
EXPECTED_MD5 = "024ccba9d224cfccf2b8835641f9812c"
SHEETS = {
    "Fig.3a": ("davis", "mse"),
    "Fig.3b": ("kiba", "mse"),
    "Fig.3c": ("davis", "pearson"),
    "Fig.3d": ("kiba", "pearson"),
}


def summary(values):
    return {
        "mean": statistics.mean(values),
        "sd_across_folds": statistics.stdev(values),
        "median": statistics.median(values),
        "per_fold_in_source_order": values,
    }


def main():
    archive_md5 = hashlib.md5(ARCHIVE.read_bytes()).hexdigest()
    if archive_md5 != EXPECTED_MD5:
        raise RuntimeError("Archive differs from the author's published Figshare checksum")
    with ZipFile(ARCHIVE) as archive:
        workbook_bytes = archive.read(MEMBER)
    formulas = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    cached = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    published = {}
    for sheet_name, (dataset, metric) in SHEETS.items():
        sheet = formulas[sheet_name]
        for setting, first_column in (("warm", 2), ("drug_cold", 10), ("target_cold", 18)):
            models = published.setdefault(dataset, {}).setdefault(setting, {})
            for column in range(first_column, first_column + 5):
                model = sheet.cell(4, column).value
                assert model in {"DeepDTA", "MONN", "BACPI", "GraphDTA", "DTIAM"}
                values = [float(sheet.cell(row, column).value) for row in range(5, 10)]
                result = summary(values)
                cached_mean = cached[sheet_name].cell(10, column).value
                if cached_mean is not None:
                    assert abs(float(cached_mean) - result["mean"]) < 1e-10
                result["source_cells"] = f"{sheet_name}!{sheet.cell(5, column).coordinate}:{sheet.cell(9, column).coordinate}"
                result["source_mean_formula"] = sheet.cell(10, column).value
                models.setdefault(model, {})[metric] = result

    ours = json.loads((PROJECT / "reports/ijms_readiness_audit_20261016.json").read_text())["five_fold_metrics"]
    reference = published["davis"]["drug_cold"]["DTIAM"]
    comparisons = {}
    for family in ("legacy_gated", "scratch_gated", "transfer_gated"):
        comparisons[family] = {
            "mse_mean": ours[family]["mse"]["mean"],
            "pearson_mean": ours[family]["pearson"]["mean"],
            "relative_mse_reduction_against_published_dtiam": 1 - ours[family]["mse"]["mean"] / reference["mse"]["mean"],
            "pearson_gain_against_published_dtiam": ours[family]["pearson"]["mean"] - reference["pearson"]["mean"],
        }
    output = {
        "source": {
            "doi": "https://doi.org/10.6084/m9.figshare.27896790.v1",
            "article": "https://www.nature.com/articles/s41467-025-57828-0",
            "download_url": "https://ndownloader.figshare.com/files/51785909",
            "archive_md5": archive_md5,
            "workbook_member": MEMBER,
            "workbook_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
            "retrieved_on": "2026-10-16",
        },
        "published_figure3_results": published,
        "davis_drug_cold_comparison": comparisons,
        "interpretation": {
            "status": "Numerical comparison to author-published results, not a local DTIAM reproduction or a significance claim.",
            "protocol": "Our full Davis outer splits reproduce the official code's five-fold drug-cold splitter (seed 0). Our inner entity-disjoint validation differs from the official AutoGluon training path.",
            "correction": "Earlier visual estimate of DTIAM median MSE around 0.55 was inaccurate. Source mean is 0.5983246 and median is 0.608079.",
            "scope": "Observed advantage concerns Davis drug-cold only. KIBA and other settings require our completed results before comparison.",
        },
    }
    path = Path(__file__).with_name("dtiam_published_comparison_20261016.json")
    path.write_text(json.dumps(output, indent=2) + "\n")
    print(json.dumps({"DTIAM_davis_drug_cold": reference, "comparisons": comparisons}, indent=2))
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
