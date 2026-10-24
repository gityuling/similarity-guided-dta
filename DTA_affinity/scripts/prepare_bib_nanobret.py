#!/usr/bin/env python3
"""Audit the complete 2026 cellular panels before frozen-model evaluation."""

import json
from collections import defaultdict

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
from rdkit import Chem, RDLogger, rdBase
from rdkit.Chem.MolStandardize import rdMolStandardize

from bib_common import PROJECT, ROOT, canonical, record, sha256, write_once
from fetch_bib_nanobret_sources import retrieve


OUT = PROJECT / "reports/bib_nanobret_20261018"
SOURCE = OUT / "sources/PMC12951567"
KIBA = PROJECT / "data/processed/kiba_dtiam_protocol/table.csv"
DAVIS = PROJECT / "data/processed/davis_dtiam_protocol/table.csv"
# sheet, compound, endpoint, first data row, target/value/replicate columns (zero based)
SHEETS = [
    ("Dasatinib K240 DR", "Dasatinib", "ic50_nm", 4, 0, 1, None),
    ("Sorafenib K240 DR", "Sorafenib", "ic50_nm", 4, 0, 1, None),
    ("AST487 K240 DR", "AST487", "ic50_nm", 4, 0, 3, None),
    ("Foretinib K300 DR", "Foretinib", "ic50_nm", 2, 0, 3, None),
    ("LXH254 K240 DR", "LXH254", "ic50_nm", 4, 0, 3, None),
    ("AMG-925, 1 uM", "AMG925", "occupancy_pct", 2, 0, 3, (1, 2)),
    ("Ponatinib K300 DR", "Ponatinib", "ic50_nm", 2, 0, 3, None),
    ("GW296115X, 1 uM", "GW296115X", "occupancy_pct", 2, 1, 6, (4, 5)),
    ("GW830263A, 1 uM", "GW830263A", "occupancy_pct", 2, 1, 6, (4, 5)),
    ("TPKI-39, 1 uM", "TPKI-39", "occupancy_pct", 2, 1, 6, (4, 5)),
]


def save_json(name, payload):
    write_once(OUT / name, json.dumps(payload, indent=2, ensure_ascii=False, allow_nan=False) + "\n")


def parent_key(smiles):
    mol = rdMolStandardize.FragmentParent(Chem.MolFromSmiles(smiles))
    mol = rdMolStandardize.Uncharger().uncharge(mol)
    key = Chem.MolToInchiKey(mol)
    if not key:
        raise ValueError("Cannot assign standardized parent InChIKey")
    return key.split("-")[0]


def uniprot_annotations(proteins):
    folder = OUT / "sources/uniprot"
    accessions = sorted(proteins.protein_id.unique())
    sources, frames = [], []
    for start in range(0, len(accessions), 40):
        path = folder / f"kiba_{start // 40:02d}.tsv"
        item = retrieve("https://rest.uniprot.org/uniprotkb/search", path, {
            "query": " OR ".join(f"accession:{x}" for x in accessions[start:start + 40]),
            "format": "tsv", "fields": "accession,gene_names,organism_id,sequence", "size": 500,
        })
        sources.append(item)
        frames.append(pd.read_csv(path, sep="\t"))
    manifest = folder / "download_manifest.json"
    if not manifest.exists():
        write_once(manifest, json.dumps(sources, indent=2) + "\n")
    annotations = pd.concat(frames, ignore_index=True).drop_duplicates()
    if annotations.Entry.duplicated().any():
        raise RuntimeError("Conflicting UniProt records")
    joined = proteins.merge(annotations, left_on="protein_id", right_on="Entry", how="left", validate="one_to_one")
    joined["canonical_sequence_matches"] = joined.Sequence_x.eq(joined.Sequence_y)
    joined = joined.rename(columns={"Sequence_x": "Sequence", "Sequence_y": "current_uniprot_sequence"})
    write_once(OUT / "protein_annotation_audit.csv", joined.to_csv(index=False))
    aliases = defaultdict(set)
    for row in joined.to_dict("records"):
        if row.get("Organism (ID)") == 9606:
            for alias in str(row["Gene Names"]).split():
                aliases[alias.upper()].add(row["protein_id"])
    return joined.set_index("protein_id"), aliases


def parse_workbook():
    path = SOURCE / "jm5c02916_si_004.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if set(workbook.sheetnames) != {x[0] for x in SHEETS}:
        raise RuntimeError("Source workbook sheet inventory changed")
    records, sheet_checks = [], []
    for sheet, compound, endpoint, first, target_col, value_col, replicates in SHEETS:
        count, blank = 0, 0
        for rownum, values in enumerate(workbook[sheet].values, 1):
            if rownum < first:
                continue
            name, raw = values[target_col], values[value_col]
            if name is None and raw is None:
                blank += 1
                continue
            if name is None or raw is None:
                raise RuntimeError(f"Partial measurement at {sheet}:{rownum}")
            text = str(raw).strip()
            qualifier = text[0] if text.startswith((">", "<")) else "="
            value = float(text[1:] if qualifier != "=" else text)
            if not np.isfinite(value) or (endpoint == "ic50_nm" and value <= 0):
                raise RuntimeError(f"Invalid measurement at {sheet}:{rownum}")
            replicate_values = [float(values[j]) for j in replicates] if replicates else []
            if replicates and not np.isclose(np.mean(replicate_values), value, atol=1e-6, rtol=1e-8):
                raise RuntimeError(f"Replicate mean mismatch at {sheet}:{rownum}")
            k240_boundary = endpoint == "ic50_nm" and "K240" in sheet and value in (3.0, 10000.0)
            count += 1
            records.append({
                "measurement_id": f"nb2026_{len(records):04d}", "compound": compound,
                "source_sheet": sheet, "source_row": rownum,
                "source_target_cell": f"{get_column_letter(target_col + 1)}{rownum}",
                "source_value_cell": f"{get_column_letter(value_col + 1)}{rownum}",
                "target_name": str(name).strip(), "endpoint": endpoint, "raw_value": text,
                "qualifier": qualifier, "value": value,
                "endpoint_score": 9 - np.log10(value) if endpoint == "ic50_nm" else value,
                "active": value < 1000 if endpoint == "ic50_nm" else value > 50,
                "concentration_um": 1 if endpoint == "occupancy_pct" else None,
                "reported_censor": qualifier != "=", "possible_assay_boundary": k240_boundary,
                "occupancy_outside_0_100": endpoint == "occupancy_pct" and not 0 <= value <= 100,
                "replicate_1": replicate_values[0] if replicates else None,
                "replicate_2": replicate_values[1] if replicates else None,
            })
        sheet_checks.append({"sheet": sheet, "measurements": count, "formatted_blank_rows": blank})
    workbook.close()
    return pd.DataFrame(records), sheet_checks


def main():
    RDLogger.DisableLog("rdApp.warning")
    for pmcid in ("PMC5814754", "PMC12951567"):
        for item in json.loads((OUT / "sources" / pmcid / "download_manifest.json").read_text())["files"]:
            if sha256(ROOT / item["path"]) != item["sha256"]:
                raise RuntimeError("Downloaded source changed")
    table = pd.read_csv(KIBA)
    proteins, aliases = uniprot_annotations(table[["protein_id", "Sequence"]].drop_duplicates())
    compounds = pd.read_csv(SOURCE / "jm5c02916_si_002.csv").set_index("Compound name")
    maps = {}
    for name, path in (("kiba", KIBA), ("davis", DAVIS)):
        drugs = pd.read_csv(path, usecols=["drug_id", "Canonical_SMILES"]).drop_duplicates()
        drugs["identity"] = drugs.Canonical_SMILES.map(canonical)
        drugs["parent_key"] = drugs.Canonical_SMILES.map(parent_key)
        maps[name] = drugs
    manifests = {f: pd.read_csv(PROJECT / f"data/processed/kiba_identity_20261017/folds/drug_coldstart/fold_{f}.csv") for f in range(5)}
    fold_drugs = {f: table[["row_id", "drug_id", "Sequence"]].merge(m, on="row_id", validate="one_to_one") for f, m in manifests.items()}
    compound_rows = []
    for compound in sorted({x[1] for x in SHEETS}):
        source_smiles = compounds.loc[compound, "SMILES"]
        identity, parent = canonical(source_smiles), parent_key(source_smiles)
        row = {"compound": compound, "source_smiles": source_smiles,
               "Canonical_SMILES": identity, "parent_connectivity_key": parent}
        for name, drugs in maps.items():
            exact = drugs.loc[drugs.identity.eq(identity), "drug_id"].astype(str).tolist()
            conservative = drugs.loc[drugs.parent_key.eq(parent), "drug_id"].astype(str).tolist()
            row[f"{name}_exact_ids"] = "|".join(sorted(exact))
            row[f"{name}_parent_ids"] = "|".join(sorted(conservative))
        overlapping_ids = set(row["kiba_exact_ids"].split("|") + row["kiba_parent_ids"].split("|")) - {""}
        eligible = []
        for f, split in fold_drugs.items():
            seen = split.loc[split.split.isin(["train", "val"]), "drug_id"].astype(str)
            if not overlapping_ids.intersection(seen):
                eligible.append(f)
        row["eligible_folds"] = "|".join(map(str, eligible))
        row["kiba_external"] = not bool(overlapping_ids)
        row["davis_external"] = not bool(row["davis_exact_ids"] or row["davis_parent_ids"])
        compound_rows.append(row)
    compound_audit = pd.DataFrame(compound_rows)
    write_once(OUT / "compound_audit.csv", compound_audit.to_csv(index=False))
    measurements, sheets = parse_workbook()
    duplicate = measurements.duplicated(["compound", "target_name"], keep=False)
    mapped = []
    for index, row in measurements.iterrows():
        name = row.target_name
        hits = sorted(aliases.get(name.upper(), set()))
        reason = ""
        if duplicate.loc[index]:
            reason = "duplicate_source_target_ambiguous_construct"
        elif "+" in name or "cyclin" in name.lower() or "(" in name:
            reason = "complex_or_modified_construct"
        elif len(hits) != 1:
            reason = "no_unique_kiba_uniprot_gene_alias"
        elif not proteins.loc[hits[0], "canonical_sequence_matches"]:
            reason = "kiba_sequence_differs_from_current_canonical"
        compound = compound_audit.set_index("compound").loc[row.compound]
        if not reason and not compound.eligible_folds:
            reason = "no_model_free_of_train_validation_overlap"
        protein = hits[0] if len(hits) == 1 else None
        if not reason:
            sequence = proteins.loc[protein, "Sequence"]
            for f in map(int, compound.eligible_folds.split("|")):
                if not ((fold_drugs[f].split == "train") & (fold_drugs[f].Sequence == sequence)).any():
                    reason = "no_same_target_training_support"
        mapped.append({"protein_id": protein, "exclusion_reason": reason})
    measurements = pd.concat([measurements, pd.DataFrame(mapped)], axis=1)
    # Two aliases resolving to the same accession would otherwise double-count one target.
    eligible = measurements.exclusion_reason.eq("")
    mapped_duplicates = measurements.loc[eligible].duplicated(["compound", "protein_id"], keep=False)
    measurements.loc[mapped_duplicates.index[mapped_duplicates], "exclusion_reason"] = "duplicate_mapped_accession"
    measurements["included"] = measurements.exclusion_reason.eq("")
    write_once(OUT / "measurements.csv", measurements.to_csv(index=False))
    query = measurements.loc[measurements.included, ["measurement_id", "compound", "protein_id", "endpoint"]].copy()
    query = query.merge(compound_audit[["compound", "Canonical_SMILES", "eligible_folds", "kiba_external", "davis_external"]], on="compound", validate="many_to_one")
    query = query.merge(proteins[["Sequence"]], left_on="protein_id", right_index=True, validate="many_to_one")
    query.insert(0, "row_id", np.arange(len(query)))
    if query.duplicated(["compound", "protein_id"]).any():
        raise RuntimeError("Repeated query")
    write_once(OUT / "query_table.csv", query.to_csv(index=False))
    coverage = measurements.groupby(["compound", "endpoint"], sort=True).agg(
        measured_pairs=("measurement_id", "size"), included_pairs=("included", "sum"),
        outside_occupancy_range=("occupancy_outside_0_100", "sum"),
        explicit_censors=("reported_censor", "sum"), possible_boundary_values=("possible_assay_boundary", "sum"),
    ).reset_index()
    coverage = coverage.merge(compound_audit[["compound", "kiba_external", "davis_external", "eligible_folds"]], on="compound", validate="one_to_one")
    write_once(OUT / "coverage.csv", coverage.to_csv(index=False))
    plan = {
        "study": "PMC12951567 complete SI004 panels; source chosen for coverage before model scoring",
        "design": "Exploratory external cellular validation after the negative BRAF/A375 viability case",
        "model_dataset": "KIBA; existing identity-disjoint five-fold models, seed 20261012",
        "model_groups": ["deep", "neighbor_prior", "fixed_blend", "similarity_gate"],
        "no_training_or_new_endpoint_fitting": True,
        "primary": "Within-compound Spearman with higher engagement endpoint_score; macro-average compounds, IC50 and occupancy separately",
        "primary_contrast": "similarity_gate minus deep; fixed_blend and neighbor_prior contrasts retained",
        "minimum_pairs_per_compound": 20,
        "secondary": "AUROC, AP, top10 hits and EF10 per compound; active IC50 <1000 nM or occupancy >50% at 1 uM",
        "endpoint_processing": "IC50: 9-log10(value nM); reported inequalities represented by boundary with tied ranks. K240 numeric 3/10000 boundaries flagged, not assumed exact uncensored affinities. Occupancy: finite published mean, including negative/out-of-range values, no primary clipping.",
        "duplicate_policy": "Exclude all repeated compound-target source names and repeated mapped accessions; never select the best replicate",
        "target_policy": "Unique exact UniProt human gene-name/synonym mapping to KIBA accession; KIBA sequence must match current canonical. Exclude complexes/modified-name constructs. NanoLuc fusion and experimental isoform sequence not independently verified.",
        "overlap_policy": "RDKit isomeric canonical identity plus uncharged largest-parent standard InChI connectivity block; stereo-insensitive conservative overlap. Every used fold excludes matching IDs from BOTH train and validation; average only eligible folds.",
        "strata": "Report all eligible compounds and the Davis-disjoint subset; KIBA-new and heldout-known drugs separately if both exist. Davis overlap is development exposure, not KIBA training exposure.",
        "sensitivities": ["occupancy clipped to [0,100]", "IC50 exclude explicit inequalities and flagged K240 3/10000 bounds"],
        "uncertainty": "2000 paired compound-cluster bootstrap draws of macro metric difference, seed 20261019; conditional on fixed models and sampled panels, very few compounds, no multiplicity adjustment; not independent-target bootstrap",
        "bootstrap_n": 2000, "bootstrap_seed": 20261019,
        "limitations": ["No KIBA-score versus cellular-IC50 MSE", "Cellular engagement does not validate cell killing or disease efficacy", "Generalization limited to measured known kinase targets", "No independent DTIAM predictions for this panel", "Pretrained encoder exposure cannot be excluded"],
        "rdkit_version": rdBase.rdkitVersion,
    }
    save_json("plan.json", plan)
    audit = {"raw_measurements": len(measurements), "included_pairs": len(query),
             "included_compounds": query.compound.nunique(), "included_targets": query.protein_id.nunique(),
             "exclusions": measurements.loc[~measurements.included].exclusion_reason.value_counts().to_dict(),
             "source_sheet_checks": sheets, "compound_coverage": coverage.to_dict("records"),
             "protein_canonical_sequence_matches": int(proteins.canonical_sequence_matches.sum()),
             "protein_accessions": len(proteins)}
    save_json("audit.json", audit)
    inputs = [KIBA, DAVIS, SOURCE / "jm5c02916_si_004.xlsx", SOURCE / "jm5c02916_si_002.csv",
              SOURCE / "article.xml", SOURCE / "jm5c02916_si_005.csv", PROJECT / "scripts/prepare_bib_nanobret.py"]
    inputs += [OUT / name for name in ("plan.json", "audit.json", "measurements.csv", "query_table.csv", "compound_audit.csv", "protein_annotation_audit.csv", "coverage.csv")]
    inputs += sorted((OUT / "sources/uniprot").glob("*.tsv"))
    for f in range(5):
        stem = f"bib_kiba_identity_f{f}_seed20261012"
        inputs += [PROJECT / f"results/checkpoints/nexusdti_{stem}.pt", PROJECT / f"results/{stem}_ablation.json",
                   PROJECT / f"results/bib_kiba_identity_f{f}_frozen_prior.json",
                   PROJECT / f"data/processed/kiba_identity_20261017/folds/drug_coldstart/fold_{f}.csv"]
    save_json("freeze.json", {"files": [record(p) for p in inputs], "model_endpoint_metrics_calculated": False})
    print(json.dumps(audit, indent=2), flush=True)


if __name__ == "__main__":
    main()
